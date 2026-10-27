# tests/unit/experiment/test_export_survey_results_cli.py

# ruff: noqa: F811, S101

"""The export entry point end to end: a batch tree in, a workbook and cards out.

The layout and the signals are pinned by their own suites; what these add is the
wiring -- that one command goes from a batch config to the three files, that a
finished ``summary.json`` goes through the same path, and that a bad input is an
error rather than an empty delivery.
"""

import importlib
import json
import sys
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

import openpyxl  # noqa: E402
from _survey_summary_helpers import (  # noqa: E402, F401
    config_for,
    load_protocol,
    manifest,
    summary_cli,
    write_config,
    write_tree,
)

pytestmark = pytest.mark.unit

_SCRIPTS = Path(__file__).resolve().parents[3] / "scripts"


@pytest.fixture(scope="module")
def export_cli():
    if str(_SCRIPTS) not in sys.path:
        sys.path.insert(0, str(_SCRIPTS))
    return importlib.import_module("export_survey_results")


def _grid():
    c_tokens = [float(token) for token in load_protocol()["c_tokens"]["scar"]]
    return {
        f"c_{c}_seed_{seed}": manifest(c=c, seed=seed, accuracy=0.8 + 0.01 * seed)
        for c in c_tokens
        for seed in range(5)
    }


def _config(tmp_path):
    tree = write_tree(tmp_path / "tree", _grid())
    return write_config(tmp_path / "roots.json", config_for({"B1": (tree, len(_grid()))}))


def test_basic_one_command_goes_from_a_batch_config_to_the_three_deliverables(
    export_cli, tmp_path, capsys
):
    out = tmp_path / "out"

    code = export_cli.main(
        ["--config", str(_config(tmp_path)), "--out-dir", str(out), "--protocol", "survey-v2.0"]
    )

    assert code == 0, capsys.readouterr().err
    assert (out / "survey_results.xlsx").is_file()
    assert (out / "diagnostic_cards.md").is_file()
    cards = json.loads((out / "diagnostic_cards.json").read_text(encoding="utf-8"))
    assert "nnpu" in cards
    assert (out / "summary.json").is_file()


def test_basic_the_workbook_it_writes_reads_back_with_the_four_group_sheets(export_cli, tmp_path):
    out = tmp_path / "out"
    export_cli.main(
        ["--config", str(_config(tmp_path)), "--out-dir", str(out), "--protocol", "survey-v2.0"]
    )

    workbook = openpyxl.load_workbook(out / "survey_results.xlsx")

    assert {"SCAR-PA", "SCAR-OA", "SAR-OA", "PN-oracle", "coverage", "readme"} <= set(
        workbook.sheetnames
    )
    flat = [c for r in workbook["SCAR-OA"].iter_rows(values_only=True) for c in r if c]
    assert "nnpu" in flat
    assert any(isinstance(c, str) and "±" in c for c in flat)


def test_basic_a_finished_summary_json_exports_the_same_way(export_cli, summary_cli, tmp_path):
    summary_dir = tmp_path / "summary"
    summary_cli.main(
        [
            "--config",
            str(_config(tmp_path)),
            "--out-dir",
            str(summary_dir),
            "--protocol",
            "survey-v2.0",
        ]
    )
    out = tmp_path / "out"

    code = export_cli.main(["--summary", str(summary_dir / "summary.json"), "--out-dir", str(out)])

    assert code == 0
    assert (out / "survey_results.xlsx").is_file()
    assert not (out / "summary.json").exists()


def test_param_exactly_one_input_is_required(export_cli, tmp_path):
    with pytest.raises(SystemExit):
        export_cli.main(["--out-dir", str(tmp_path / "out")])


def test_param_a_missing_summary_file_is_an_error_and_writes_nothing(export_cli, tmp_path, capsys):
    out = tmp_path / "out"

    code = export_cli.main(["--summary", str(tmp_path / "nope.json"), "--out-dir", str(out)])

    assert code == 1
    assert "error:" in capsys.readouterr().err
    assert not (out / "survey_results.xlsx").exists()


def test_edge_a_summary_with_an_unknown_mechanism_is_refused_not_dropped(
    export_cli, tmp_path, capsys
):
    bad = {
        "rows": [
            {
                "result_identity": {"dataset": "d", "training_path": "p", "method": "m"},
                "row_key": {
                    "comparability_group": "g",
                    "run_view": "os-compatible",
                    "mechanism": "mystery",
                    "method": "m",
                    "selection_protocol": "OA",
                    "c_token": "0.1",
                },
                "metric": {"mean": 0.5, "std": 0.1, "mean_percent": 50.0, "std_percent": 10.0},
                "status": "formal",
                "reasons": [],
            }
        ]
    }
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(bad), encoding="utf-8")

    code = export_cli.main(["--summary", str(path), "--out-dir", str(tmp_path / "out")])

    assert code == 1
    assert "mechanism" in capsys.readouterr().err


def test_determ_exporting_the_same_summary_twice_gives_identical_cards(
    export_cli, summary_cli, tmp_path
):
    summary_dir = tmp_path / "summary"
    summary_cli.main(
        [
            "--config",
            str(_config(tmp_path)),
            "--out-dir",
            str(summary_dir),
            "--protocol",
            "survey-v2.0",
        ]
    )
    summary = str(summary_dir / "summary.json")
    export_cli.main(["--summary", summary, "--out-dir", str(tmp_path / "a")])
    export_cli.main(["--summary", summary, "--out-dir", str(tmp_path / "b")])
    for name in ("diagnostic_cards.json", "diagnostic_cards.md"):
        assert (tmp_path / "a" / name).read_bytes() == (tmp_path / "b" / name).read_bytes()

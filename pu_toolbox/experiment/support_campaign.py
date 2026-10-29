"""Explicit D53 correction attempts, kept separate from baseline observations."""

from __future__ import annotations

import copy
import csv
import hashlib
import json
import subprocess
from collections import defaultdict
from pathlib import Path

from . import batch_evidence
from .pilot_plan import manifest_identity
from .survey_protocol import digest, load_protocol
from .survey_recipe_binding import SCHEMA_VERSION, bind_recipe, matches_recipe_binding
from .survey_recipe_registry import canonical_digest, load_registry


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_support_campaign(path):
    """Verify a reviewed plan and its CPU preflights before touching a controller."""
    path = Path(path).resolve()
    raw = path.read_bytes()
    plan = json.loads(raw)
    if plan.get("schema_version") != "survey-support-repair-plan-1" or plan.get("launched"):
        raise ValueError("expected an unlaunched D53 repair plan")
    source = Path(plan["source"]["path"]).resolve()
    if subprocess.check_output(["git", "status", "--porcelain"], cwd=source).strip():
        raise ValueError("repair training source is dirty")
    commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=source, text=True).strip()
    if commit != plan["source"]["code_commit"]:
        raise ValueError("repair training commit changed")
    recorder = file_hash(source / "pu_toolbox/experiment/support_checkpoints.py")
    if recorder != plan["source"]["support_recorder_sha256"] or recorder != file_hash(
        Path(__file__).with_name("support_checkpoints.py")
    ):
        raise ValueError("repair source does not contain the inspected all-epoch recorder")
    protocol = load_protocol(source / "pu_toolbox/experiment/survey_protocol_v3.json")
    if digest(protocol) != plan["protocol_sha256"]:
        raise ValueError("repair protocol changed")
    registry = load_registry(source / "pu_toolbox/experiment/survey_recipe_registry_v1.json")
    campaign_id = "d53_" + hashlib.sha256(raw).hexdigest()[:16]
    commands = plan["commands"]
    roots = {Path(c["results_root"]).resolve().parent for c in commands}
    if len(roots) != 1 or not commands:
        raise ValueError("repair plan needs one new campaign root")
    root = roots.pop()
    if root.exists():
        raise ValueError(
            "repair campaign root already exists; use its existing queue, not another mount"
        )
    if any(
        root == Path(r).resolve()
        or root.is_relative_to(Path(r).resolve())
        or Path(r).resolve().is_relative_to(root)
        for r in plan["result_roots"]
    ):
        raise ValueError("repair output overlaps historical results")
    by_unit = defaultdict(list)
    for fit in plan["fits"]:
        by_unit[(fit["dataset"], fit["profile"], fit["training_path"])].append(fit)
        for ref in fit.get("sources", []):
            if file_hash(ref["path"]) != ref["sha256"]:
                raise ValueError("original repair evidence changed")
    jobs = []
    for command in commands:
        unit = command["unit"]
        dataset, method, training_path = (unit[k] for k in ("dataset", "profile", "training_path"))
        if method not in {"lagam", "self_pu_full_meta"}:
            raise ValueError("D53 repair mount supports only the two affected profiles")
        label = f"{dataset}__{method}__{training_path}"
        preflight = json.loads((path.parent / (label + ".plan.json")).read_text())
        if preflight["source_protocol_sha256"] != plan["protocol_sha256"]:
            raise ValueError("CPU preflight is not bound to this protocol")
        wanted = {(s, c) for s in protocol["seeds"] for c in protocol["c_tokens"]["scar"]}
        cells = [(r["seed"], r["c_token"]) for r in preflight["runs"]]
        if set(cells) != wanted or len(cells) != len(wanted):
            raise ValueError("CPU preflight grid is incomplete or duplicated")
        fits = by_unit[dataset, method, training_path]
        if len(fits) != len(wanted) or any(f["status"] == "conflict" for f in fits):
            raise ValueError("repair source unit is incomplete or conflicted")
        argv = command["pilot_argv"]
        support = Path(argv[argv.index("--support-root") + 1]).resolve() / dataset
        binding = bind_recipe(
            registry,
            protocol,
            method=method,
            dataset=dataset,
            training_path=training_path,
            candidates=protocol["candidate_pool"],
            constructor_overrides={},
            binding_schema_version=SCHEMA_VERSION,
        )
        for run in preflight["runs"]:
            if (run["dataset"], run["method"], run["training_path"], run["mechanism"]) != (
                dataset,
                method,
                training_path,
                "scar",
            ):
                raise ValueError("CPU preflight unit differs from repair request")
            seed, token = run["seed"], run["c_token"]
            if method == "self_pu_full_meta" and run["view"] != "os-compatible":
                raise ValueError("full-meta correction must use the approved OS path")
            receipt_path = support / f"split_{seed}" / "support_manifest.json"
            receipt = json.loads(receipt_path.read_text())
            if receipt["dataset"] != dataset or receipt["seed"] != seed:
                raise ValueError("support receipt identity mismatch")
            canary = seed == min(protocol["seeds"]) and token == protocol["c_tokens"]["scar"][0]
            fit_label = f"{campaign_id}__{label}__s{seed}_c{token}"
            jobs.append(
                dict(
                    label=fit_label,
                    campaign_id=campaign_id,
                    campaign_root=str(root),
                    campaign_unit=label,
                    campaign_canary=canary,
                    gate_passed=False,
                    parent_label=label,
                    method=method,
                    dataset=dataset,
                    path=training_path,
                    oracle=False,
                    mechanism="scar",
                    seeds=[seed],
                    c_tokens=[token],
                    expected_runs=1,
                    expected_keys=[[dataset, method, training_path, "scar", token, seed]],
                    expected_recipe=binding,
                    expected_view=run["view"],
                    protocol_sha256=digest(protocol),
                    root=str(root / label / f"fit_s{seed}_c{token}"),
                    source=str(source),
                    source_code_commit=commit,
                    recorder_sha256=recorder,
                    dependency_lock_sha256=file_hash(source / "uv.lock"),
                    support_receipt_sha256=file_hash(receipt_path),
                    support_receipt=str(receipt_path),
                    data_root=str(support),
                    support_data=str(support),
                    class_prior=receipt["class_prior"]["population"],
                    os_or_ts="os" if method == "self_pu_full_meta" else run["view"].split("-")[0],
                    lane="gpu",
                    priority=-100,
                    status="pending" if canary else "held",
                    hold_reason=None if canary else "awaiting_full_budget_canary",
                )
            )
    keys = [tuple(j["expected_keys"][0]) for j in jobs]
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate repair fit")
    meta = dict(
        campaign_id=campaign_id,
        root=str(root),
        repair_plan=str(path),
        repair_plan_sha256=hashlib.sha256(raw).hexdigest(),
        source_code_commit=commit,
        expected_fits=len(jobs),
        gate_policy="full_budget_canary_per_unit_then_every_fit",
        recipe_registry_sha256=canonical_digest(registry),
    )
    return meta, jobs


def mount_campaign(existing, metadata, additions):
    """Idempotent mount, with intentional correction attempts namespaced explicitly."""
    current = [j for j in existing if j.get("campaign_id") == metadata["campaign_id"]]
    if current:
        if {j["label"]: j["root"] for j in current} != {j["label"]: j["root"] for j in additions}:
            raise ValueError("existing campaign differs from requested mount")
        return existing
    from .queue_scheduling import validate_fit_ownership

    combined = [*existing, *copy.deepcopy(additions)]
    validate_fit_ownership(combined)
    return combined


def release_canary_followups(jobs):
    """A failed or unverifiable sentinel blocks its 14 siblings automatically."""
    canaries = {(j["campaign_id"], j["campaign_unit"]): j for j in jobs if j.get("campaign_canary")}
    failed_units = {
        (j["campaign_id"], j["campaign_unit"])
        for j in jobs
        if j.get("campaign_id") and j["status"] in {"failed", "blocked"}
    }
    for job in jobs:
        if job["status"] not in {"held", "pending"} or not job.get("campaign_id"):
            continue
        if (job["campaign_id"], job["campaign_unit"]) in failed_units:
            job.update(status="blocked", hold_reason="unit_quality_gate_failed")
            continue
        if job["status"] != "held":
            continue
        sentinel = canaries[job["campaign_id"], job["campaign_unit"]]
        if sentinel["status"] == "completed" and sentinel.get("gate_passed"):
            job.update(status="pending", hold_reason=None)
        elif sentinel["status"] in {"failed", "blocked"}:
            job.update(status="blocked", hold_reason="unit_canary_failed")


def repair_manifest_findings(job, payload, *, checkpoint_check):
    """Fail closed on the original regression and on actual D52/supervision gaps."""
    problems = list(
        checkpoint_check(
            payload,
            load_protocol(Path(job["source"]) / "pu_toolbox/experiment/survey_protocol_v3.json"),
        )
    )
    if manifest_identity(payload) != tuple(job["expected_keys"][0]):
        problems.append("repair_fit_identity_mismatch")
    if not matches_recipe_binding(payload, job["expected_recipe"]):
        problems.append("repair_recipe_mismatch")
    if any(
        set(payload.get(field, {})) != {"OA", "PA+clean-support"}
        for field in ("selection", "test_results")
    ):
        problems.append("repair_comparison_track_mismatch")
    support = payload.get("extra_supervision", {})
    if not support.get("all_role_isolation_verified") or not support.get(
        "support_removed_from_train"
    ):
        problems.append("repair_role_isolation_unverified")
    try:
        block = payload["batch"]
        packet_path = Path(job["campaign_root"]) / "batch_evidence" / block["evidence_file"]
        if Path(block["evidence_file"]).name != block["evidence_file"]:
            raise ValueError("unsafe packet reference")
        packet = batch_evidence.load(packet_path)
        batch_evidence.require_matches(
            packet,
            unit=dict(
                dataset=job["dataset"],
                method=job["method"],
                training_path=job["path"],
                mechanism="scar",
            ),
            seeds=job["seeds"],
            c_tokens=job["c_tokens"],
            protocol_sha256=job["protocol_sha256"],
        )
        if batch_evidence.manifest_block(packet, packet_path) != block:
            raise ValueError("packet reference mismatch")
        if packet["code_commit"] != job["source_code_commit"] or packet["code_commit_dirty"]:
            raise ValueError("wrong code identity")
        if packet["dependency_lock"]["sha256"] != job["dependency_lock_sha256"]:
            raise ValueError("wrong dependency lock")
        environment = packet["environment"]
        if not environment.get("hardware", {}).get("processor") or not environment.get(
            "gpu_driver_version"
        ):
            raise ValueError("missing hardware or driver identity")
        if any(
            environment.get("thread_environment", {}).get(k) != "2"
            for k in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS")
        ):
            raise ValueError("thread identity mismatch")
    except (OSError, ValueError, KeyError, TypeError):
        problems.append("repair_batch_evidence_invalid")
    return sorted(set(problems))


def write_campaign_tables(root, jobs):
    """Export correction observations separately; never rewrite manifest/sidecars."""
    from openpyxl import Workbook

    fields = [
        "method",
        "dataset",
        "seed",
        "c",
        "selection",
        "accuracy",
        "auc",
        "f1",
        "job_status",
        "quality_gate_passed",
        "manifest",
    ]
    rows = []
    for job in jobs:
        for path in Path(job["root"]).rglob("manifest.json"):
            try:
                payload = json.loads(path.read_text())
                for selection, metrics in payload.get("test_results", {}).items():
                    rows.append(
                        dict(
                            method=job["method"],
                            dataset=job["dataset"],
                            seed=payload["seed"],
                            c=payload["c_requested_token"],
                            selection=selection,
                            **{k: metrics.get(k) for k in ("accuracy", "auc", "f1")},
                            job_status=job["status"],
                            quality_gate_passed=bool(job.get("gate_passed")),
                            manifest=str(path),
                        )
                    )
            except (OSError, ValueError, KeyError, TypeError):
                continue
    root = Path(root)
    temporary = root / "per_run.csv.new"
    with temporary.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(root / "per_run.csv")
    book = Workbook()
    sheet = book.active
    sheet.title = "correction_observations"
    sheet.append(fields)
    for row in rows:
        sheet.append([row[k] for k in fields])
    status = book.create_sheet("queue_status")
    status.append(["unit", "seed", "c", "status", "canary", "quality_gate_passed"])
    for job in jobs:
        status.append(
            [
                job["campaign_unit"],
                job["seeds"][0],
                job["c_tokens"][0],
                job["status"],
                job["campaign_canary"],
                job.get("gate_passed", False),
            ]
        )
    notes = book.create_sheet("scope")
    notes.append(["Corrected engineering observations only; QA pass is not formal admission."])
    temporary = root / "survey_results.new.xlsx"
    book.save(temporary)
    temporary.replace(root / "survey_results.xlsx")

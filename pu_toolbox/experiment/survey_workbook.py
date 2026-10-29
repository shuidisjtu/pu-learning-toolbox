# pu_toolbox/experiment/survey_workbook.py

"""The delivery workbook: the summary report laid out as protocol §5.1 asks.

The exporter reads the report ``summarize_survey_results`` builds and nothing
else, so it can neither re-derive a number nor see a manifest.  Its job is
layout, and the layout carries four rules:

* **Four groups, never mixed.**  SCAR-PA, SCAR-OA, SAR-OA (LBE-A and LBE-B in
  separate blocks) and the PN oracle each get their own sheet.  ``PA`` with
  training-time clean support is listed in its own block on the SCAR-PA sheet,
  beside the pure-PU rows and not among them.
* **Nothing is ranked.**  No rank column, no total, no cross-dataset ordering;
  every grouping column is printed so a reader can see which rows were never
  compared.
* **A row the report could not complete is not a result.**  Anything short of the
  protocol's seeds, refused, or not reproducible goes to ``coverage`` with its
  missing seeds and reasons; a probe run, and a SAR run scored by ``PA`` (which the
  protocol keeps as a diagnostic log), go to ``diagnostic``.
* **A row that fits no group is an error, not an omission.**  An unknown
  mechanism or selection protocol raises rather than vanishing from a delivery.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

#: The four protocol §5.1 groups, one sheet each.
GROUP_SHEETS: tuple[str, ...] = ("SCAR-PA", "SCAR-OA", "SAR-OA", "PN-oracle")
DIAGNOSTIC_SHEET = "diagnostic"
COVERAGE_SHEET = "coverage"
README_SHEET = "readme"

#: Statuses whose numbers may stand on a result sheet.  ``partial`` keeps its
#: numbers but not the right to be ranked; everything else is explained on
#: ``coverage``, and a probe on ``diagnostic``.
_RESULT_STATUSES = frozenset({"formal", "partial"})

_SELECTION_PROTOCOLS = ("PA", "OA", "PA+clean-support")
_SAR_BLOCKS = {"sar_lbe_a": "LBE-A", "sar_lbe_b": "LBE-B"}
_SUPPORT_BLOCK = "PA + 训练期 clean support（另列，不参与纯 PU 的 PA 优势比较）"
_PURE_BLOCK = "纯 PU 的 PA"
_PARTIAL_STATUS = "partial（有数值，不得排名）"

_RESULT_COLUMNS: tuple[str, ...] = (
    "数据集",
    "训练路径",
    "比较组",
    "视图",
    "机制",
    "方法",
    "协议",
    "标记频率 c",
    "n",
    "均值%",
    "标准差%",
    "均值±标准差",
    "AUC",
    "F1",
    "Precision",
    "Recall",
    "单配置成本均值(s)",
    "调参成本合计(s)",
    "峰值显存(GiB)",
    "状态",
    "reasons",
)

_COVERAGE_COLUMNS: tuple[str, ...] = (
    "状态",
    "数据集",
    "训练路径",
    "比较组",
    "视图",
    "机制",
    "方法",
    "协议",
    "标记频率 c",
    "已观测 seed 数",
    "应有 seed 数",
    "缺失 seed",
    "reasons",
)

_README_NOTES: tuple[str, ...] = (
    "本工作簿只重排汇总报告，不重算任何数；每格均值与标准差来自报告里的 5 seed 聚合。",
    "不生成总排名：跨数据集只比较趋势，各组、各数据集、各训练路径内也不设名次列（协议 §5.1）。",
    "单 seed 或缺 seed 的读数不进结论表，记在 coverage 页并写明缺失的 seed。",
    "SAR 下 PA 仅为诊断日志，与探针运行一起放在 diagnostic 页。",
    "PA + 训练期 clean support 在 SCAR-PA 页另列一块，不参与纯 PU 的 PA 优势比较。",
    "partial 行有数值，但不得排名；formal 以报告行的 status 为准，"
    "批次是否 formal_ready 另见报告本身。",
)


class WorkbookError(ValueError):
    """The report holds a row the workbook has no honest place for."""


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise WorkbookError(
            "openpyxl is required to write the workbook: uv sync --extra report"
        ) from exc
    return openpyxl


def _percent(block: Mapping[str, Any] | None) -> str:
    if not block or block.get("mean_percent") is None:
        return "n/a"
    text = f"{block['mean_percent']:.2f}"
    if block.get("std_percent") is not None:
        text += f" ± {block['std_percent']:.2f}"
    return text


def _additional(row: Mapping[str, Any], name: str) -> str:
    block = row.get("additional_metrics", {}).get(name)
    text = _percent(block)
    if block and block.get("unavailable_seeds"):
        text += f" ({block['n_observed']}/{block['n_seeds']} seeds)"
    return text


def _validate(row: Mapping[str, Any]) -> tuple[str, str]:
    key = row["row_key"]
    mechanism, protocol = key["mechanism"], key["selection_protocol"]
    if mechanism not in ("scar", "pn_oracle", *_SAR_BLOCKS):
        raise WorkbookError(f"unknown mechanism {mechanism!r} for method {key['method']!r}")
    if protocol not in _SELECTION_PROTOCOLS:
        raise WorkbookError(f"unknown selection protocol {protocol!r} for method {key['method']!r}")
    return mechanism, protocol


def _route(row: Mapping[str, Any]) -> tuple[str, str]:
    """The ``(sheet, block)`` a row is filed under."""
    mechanism, protocol = _validate(row)
    status = row["status"]
    # A probe is a smoke run scored against the full seed grid, so it is short of
    # seeds by construction; it is a technical comparison, not an unfinished result.
    if status == "technical_probe":
        return DIAGNOSTIC_SHEET, ""
    if row["metric"].get("missing_seeds") or status not in _RESULT_STATUSES:
        return COVERAGE_SHEET, ""
    if mechanism == "scar":
        if protocol == "OA":
            return GROUP_SHEETS[1], ""
        return GROUP_SHEETS[0], _PURE_BLOCK if protocol == "PA" else _SUPPORT_BLOCK
    if mechanism == "pn_oracle":
        return (GROUP_SHEETS[3], "") if protocol == "OA" else (DIAGNOSTIC_SHEET, "")
    return (GROUP_SHEETS[2], _SAR_BLOCKS[mechanism]) if protocol == "OA" else (DIAGNOSTIC_SHEET, "")


def _sort_key(row: Mapping[str, Any]) -> tuple[str, ...]:
    key = row["row_key"]
    identity = row["result_identity"]
    return (
        identity["dataset"],
        identity["training_path"],
        key["run_view"],
        key["mechanism"],
        key["method"],
        key["selection_protocol"],
        str(key["c_token"]),
    )


#: Blocks in the order a reader meets them: pure PU before the support rows that
#: sit beside it, and LBE-A before LBE-B.
_BLOCK_ORDER = ("", _PURE_BLOCK, _SUPPORT_BLOCK, "LBE-A", "LBE-B")


def _ordered(blocks: Mapping[str, list]) -> list[tuple[str, list]]:
    """The filed blocks in reading order; an empty sheet still gets one empty block."""
    return sorted(blocks.items(), key=lambda item: _BLOCK_ORDER.index(item[0])) or [("", [])]


def _costs(report: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    return {item["cost_id"]: item["costs"] for item in report.get("run_costs", [])}


def _result_cells(row: Mapping[str, Any], costs: Mapping[str, Any]) -> list[Any]:
    key, identity, metric = row["row_key"], row["result_identity"], row["metric"]
    cost = costs.get(row["cost_id"]) or {}
    single = cost.get("single_configuration_cost_seconds")
    tuning = cost.get("tuning_cost_seconds")
    peak = cost.get("peak_gpu_memory_bytes")
    return [
        identity["dataset"],
        identity["training_path"],
        key["comparability_group"],
        key["run_view"],
        key["mechanism"],
        key["method"],
        key["selection_protocol"],
        key["c_token"],
        f"{metric['n_observed']}/{metric['n_expected']}",
        metric["mean_percent"],
        metric["std_percent"],
        _percent(metric),
        _additional(row, "auc"),
        _additional(row, "f1"),
        _additional(row, "precision"),
        _additional(row, "recall"),
        single["mean"] if single else None,
        tuning["sum"] if tuning else None,
        None if peak is None else peak / 1024**3,
        _PARTIAL_STATUS if row["status"] == "partial" else row["status"],
        ", ".join(row["reasons"]),
    ]


def _coverage_cells(row: Mapping[str, Any]) -> list[Any]:
    key, identity, metric = row["row_key"], row["result_identity"], row["metric"]
    return [
        row["status"],
        identity["dataset"],
        identity["training_path"],
        key["comparability_group"],
        key["run_view"],
        key["mechanism"],
        key["method"],
        key["selection_protocol"],
        key["c_token"],
        metric["n_observed"],
        metric["n_expected"],
        ", ".join(str(seed) for seed in metric["missing_seeds"]),
        ", ".join(row["reasons"]),
    ]


def _bold(openpyxl, cells):
    font = openpyxl.styles.Font(bold=True)
    for cell in cells:
        cell.font = font


def _write_table(ws, openpyxl, columns: Sequence[str], blocks: Sequence[tuple[str, list]]) -> None:
    """Write each ``(title, rows)`` block under its own header row."""
    for index, (title, rows) in enumerate(blocks):
        if index:
            ws.append([])
        if title:
            ws.append([title])
            _bold(openpyxl, ws[ws.max_row])
        ws.append(list(columns))
        _bold(openpyxl, ws[ws.max_row])
        for cells in rows:
            ws.append(cells)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        longest = max(
            (len(str(cell.value)) for cell in column if cell.value is not None), default=8
        )
        ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 8), 48)


def build_workbook(report: Mapping[str, Any]):
    """Lay *report* out as the delivery workbook and return it (nothing is saved)."""
    openpyxl = _require_openpyxl()
    costs = _costs(report)

    filed: dict[str, dict[str, list]] = {
        sheet: {} for sheet in (*GROUP_SHEETS, DIAGNOSTIC_SHEET, COVERAGE_SHEET)
    }
    counts: dict[str, int] = {}
    for row in sorted(report.get("rows", []), key=_sort_key):
        sheet, block = _route(row)
        cells = _coverage_cells(row) if sheet == COVERAGE_SHEET else _result_cells(row, costs)
        filed[sheet].setdefault(block, []).append(cells)
        counts[row["status"]] = counts.get(row["status"], 0) + 1

    workbook = openpyxl.Workbook()
    readme = workbook.active
    readme.title = README_SHEET
    facts = [
        ("协议版本", report.get("protocol_version")),
        ("协议 sha256", report.get("protocol_sha256")),
        ("报告生成时间", report.get("generated_at")),
        ("报告总体判定", report.get("overall")),
        ("正式准入", report.get("admission_status")),
        ("报告行数", len(report.get("rows", []))),
    ]
    facts.extend((f"状态={status}", count) for status, count in sorted(counts.items()))
    facts.extend(
        ("A11检查说明", check["message"])
        for check in report.get("checks", [])
        if check["check_id"] == "A11" and check.get("observed", {}).get("waived_batches")
    )
    for item in facts:
        readme.append(list(item))
    readme.append([])
    for note in _README_NOTES:
        readme.append([note])
    _bold(openpyxl, readme["A"])
    readme.column_dimensions["A"].width = 34
    readme.column_dimensions["B"].width = 70

    for sheet in (*GROUP_SHEETS, DIAGNOSTIC_SHEET):
        ws = workbook.create_sheet(sheet)
        _write_table(ws, openpyxl, _RESULT_COLUMNS, _ordered(filed[sheet]))
    ws = workbook.create_sheet(COVERAGE_SHEET)
    _write_table(ws, openpyxl, _COVERAGE_COLUMNS, _ordered(filed[COVERAGE_SHEET]))
    return workbook


def write_workbook(report: Mapping[str, Any], path: str | Path) -> Path:
    """Build the workbook and save it to *path*."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(report).save(target)
    return target
